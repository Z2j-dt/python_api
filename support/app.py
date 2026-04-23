# -*- coding: utf-8 -*-
"""
技术支持 - 单入口 app.py，合并 Excel 入库 Hive（可扩展）。
support 下仅 app.py + pipeline.py（内部模块），templates、static、column_map 在 support/。
"""
import os
import re
import time
from pathlib import Path
from datetime import datetime, date

import pymysql
from flask import Flask, request, jsonify, render_template, redirect
from werkzeug.middleware.dispatcher import DispatcherMiddleware

from dolphinscheduler_query_failed_tasks import (
    DolphinSchedulerAPI,
    StarRocksStore,
    DEFAULT_SR_CONF,
    sync_failed_tasks_for_projects,
)

_SUPPORT = Path(__file__).resolve().parent
ROOT = _SUPPORT.parent
if str(ROOT) not in __import__("sys").path:
    __import__("sys").path.insert(0, str(ROOT))

from portal.config import Config as PortalConfig
from portal.middleware import auth_middleware
from support.pipeline import run_pipeline, _allowed_ext

# Excel 入库配置
_EXCEL_CONFIG = {
    "TMP_DATABASE": os.environ.get("EXCEL_TMP_DATABASE", "tmp"),
    "IMPALA_HOST": os.environ.get("IMPALA_HOST", "192.168.168.219"),
    "IMPALA_PORT": int(os.environ.get("IMPALA_PORT", "21000")),
    "IMPALA_USER": os.environ.get("IMPALA_USER", "hive"),
    "IMPALA_PASSWORD": os.environ.get("IMPALA_PASSWORD", ""),
    "IMPALA_RUN_AS_OS_USER": os.environ.get("IMPALA_RUN_AS_OS_USER", "hive"),
    "HDFS_BASE_PATH": os.environ.get("HDFS_BASE_PATH", "/user/hive/warehouse/tmp.db/"),
    "HDFS_UPLOAD_AS_USER": os.environ.get("HDFS_UPLOAD_AS_USER", "tenant_sync"),
    "CSV_ENCODING": os.environ.get("CSV_ENCODING", "UTF-8"),
    "UPLOAD_FOLDER": os.environ.get("EXCEL_UPLOAD_FOLDER", "/tmp/excel_to_hive_uploads"),
    "MAX_CONTENT_LENGTH": 50 * 1024 * 1024,
}

_excel_app = Flask("excel_to_hive", template_folder=str(_SUPPORT / "templates"), static_folder=str(_SUPPORT / "static"), static_url_path="/static")
_excel_app.config.update(_EXCEL_CONFIG)
upload_dir = Path(_excel_app.config["UPLOAD_FOLDER"])
upload_dir.mkdir(parents=True, exist_ok=True)
_excel_app.config["UPLOAD_FOLDER"] = str(upload_dir)

# DolphinScheduler 重跑子应用（与 Excel 复用同一套 templates/static 目录）
_dolphin_app = Flask(
    "dolphin_support",
    template_folder=str(_SUPPORT / "templates"),
    static_folder=str(_SUPPORT / "static"),
    static_url_path="/static",
)
_sr_store = None

# SQL 导出 Excel 子应用
_sql_app = Flask(
    "sql_to_excel",
    template_folder=str(_SUPPORT / "templates"),
    static_folder=str(_SUPPORT / "static"),
    static_url_path="/static",
)


@_dolphin_app.route("/")
def dolphin_index():
    # 一个简单的前端页面，用于展示 SR 中的失败任务并操作重跑/刷新
    return render_template("dolphin_failed_tasks.html")


@_dolphin_app.route("/health")
def dolphin_health():
    return jsonify({"status": "ok", "service": "dolphin-rerun"})


@_dolphin_app.route("/api/rerun/single", methods=["POST"])
def dolphin_rerun_single():
    """
    输入: JSON { projectCode: str, processInstanceId: int }
    按流程实例重跑，调用 execute_process_instance(REPEAT_RUNNING)
    """
    data = request.get_json(force=True, silent=True) or {}
    project_code = data.get("projectCode")
    process_instance_id = data.get("processInstanceId")
    if not project_code or process_instance_id is None:
        return jsonify({"success": False, "message": "projectCode 和 processInstanceId 必填"}), 400
    try:
        client = _get_dolphin_client()
        resp = client.execute_process_instance(str(project_code), int(process_instance_id), "REPEAT_RUNNING")
        return jsonify({"success": True, "data": resp})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@_dolphin_app.route("/api/rerun/batch", methods=["POST"])
def dolphin_rerun_batch():
    """
    输入: JSON { projectCode: str, processInstanceIds: [int,...], batchSize?:5, intervalSeconds?:30 }
    按 processInstanceId 去重后逐个调用 execute_process_instance(REPEAT_RUNNING)
    """
    data = request.get_json(force=True, silent=True) or {}
    project_code = data.get("projectCode")
    process_instance_ids = data.get("processInstanceIds") or []
    batch_size = int(data.get("batchSize") or 5)
    interval_seconds = int(data.get("intervalSeconds") or 30)
    if not project_code or not process_instance_ids:
        return jsonify({"success": False, "message": "projectCode 和 processInstanceIds 必填"}), 400
    try:
        client = _get_dolphin_client()
        seen = set()
        unique_ids = []
        for pid in process_instance_ids:
            p = int(pid)
            if p not in seen:
                seen.add(p)
                unique_ids.append(p)
        results = []
        idx = 0
        while idx < len(unique_ids):
            batch = unique_ids[idx: idx + batch_size]
            for pid in batch:
                try:
                    resp = client.execute_process_instance(str(project_code), pid, "REPEAT_RUNNING")
                    results.append({"processInstanceId": pid, "success": True, "data": resp})
                except Exception as e:
                    results.append({"processInstanceId": pid, "success": False, "message": str(e)})
            idx += batch_size
            if idx < len(unique_ids):
                time.sleep(interval_seconds)
        return jsonify({"success": True, "results": results})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@_dolphin_app.route("/api/failed-tasks", methods=["GET"])
def dolphin_failed_tasks():
    """
    从 StarRocks 查询失败任务列表，用于前端展示。
    Query params:
      - date (必填): YYYY-MM-DD
      - projectCode (可选)
    """
    query_date = (request.args.get("date") or "").strip()
    project_code = (request.args.get("projectCode") or "").strip() or None
    if not query_date:
        return jsonify({"success": False, "message": "缺少必填参数 date (YYYY-MM-DD)"}), 400
    try:
        store = _get_sr_store()
        rows = store.list_failed_tasks(query_date=query_date, project_code=project_code)
        return jsonify({"success": True, "data": rows, "count": len(rows)})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@_dolphin_app.route("/api/refresh", methods=["POST"])
def dolphin_refresh():
    """
    手动刷新失败任务清单（仅刷新，不重跑）：
      - 清空 StarRocks 表
      - 按项目顺序从 DolphinScheduler 拉取指定日期的失败任务并落表
    请求 JSON:
      - date: 可选，默认当天，格式 YYYY-MM-DD
      - projectCodes: 可选，项目编码数组；不传则使用脚本里的默认顺序
    """
    data = request.get_json(force=True, silent=True) or {}
    query_date = (data.get("date") or "").strip()
    if not query_date:
        query_date = datetime.now().strftime("%Y-%m-%d")
    project_codes = data.get("projectCodes")

    # 如果前端没有传项目列表，则沿用脚本中的默认顺序（通过环境变量覆盖）。
    from dolphinscheduler_query_failed_tasks import DEFAULT_PROJECT_CODES_ORDER as _DEF_CODES
    env_codes = (os.environ.get("DOLPHIN_PROJECT_CODES") or "").strip()
    default_codes = (
        [c.strip() for c in env_codes.split(",") if c.strip()]
        if env_codes
        else _DEF_CODES
    )
    if project_codes:
        project_codes_order = [str(c).strip() for c in project_codes if str(c).strip()]
    else:
        project_codes_order = default_codes

    if not project_codes_order:
        return jsonify({"success": False, "message": "没有可用的项目编码，请检查配置或传入 projectCodes"}), 400

    try:
        client = _get_dolphin_client()
        store = _get_sr_store()
        summary = sync_failed_tasks_for_projects(
            client=client,
            store=store,
            query_date=query_date,
            project_codes_order=project_codes_order,
            auto_rerun=False,
            batch_size=5,
            interval_seconds=30,
        )
        return jsonify({"success": True, "summary": summary})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


# 服务端不做自动重跑：仅提供 刷新 SR、单条重跑、批量重跑。自动重跑接口已移除。

# DolphinScheduler 重跑配置
_DOLPHIN_CONFIG = {
    "BASE_URL": os.environ.get("DOLPHIN_BASE_URL", "http://192.168.168.219:12345/dolphinscheduler").rstrip("/"),
    "TOKEN": os.environ.get("DOLPHIN_TOKEN", ""),
}


def _get_dolphin_client() -> DolphinSchedulerAPI:
    if not _DOLPHIN_CONFIG["TOKEN"]:
        raise RuntimeError("未配置 DOLPHIN_TOKEN 环境变量，无法调用重跑接口")
    return DolphinSchedulerAPI(base_url=_DOLPHIN_CONFIG["BASE_URL"], token=_DOLPHIN_CONFIG["TOKEN"])


def _get_sr_store() -> StarRocksStore:
    global _sr_store
    if _sr_store is None:
        _sr_store = StarRocksStore(
            host=os.environ.get("STARROCKS_HOST", DEFAULT_SR_CONF["host"]),
            port=int(os.environ.get("STARROCKS_PORT", DEFAULT_SR_CONF["port"])),
            user=os.environ.get("STARROCKS_USER", DEFAULT_SR_CONF["user"]),
            password=os.environ.get("STARROCKS_PASSWORD", DEFAULT_SR_CONF["password"]),
            database=os.environ.get("STARROCKS_DATABASE", DEFAULT_SR_CONF["database"]),
            table=os.environ.get("STARROCKS_TABLE", DEFAULT_SR_CONF["table"]),
        )
    return _sr_store


def _get_hsgt_conn():
    """
    复用 StarRocks 基本连接信息，只是库名默认用 portal_db（可通过环境变量 HSGT_DATABASE 覆盖）。
    """
    return pymysql.connect(
        host=os.environ.get("STARROCKS_HOST", DEFAULT_SR_CONF["host"]),
        port=int(os.environ.get("STARROCKS_PORT", DEFAULT_SR_CONF["port"])),
        user=os.environ.get("STARROCKS_USER", DEFAULT_SR_CONF["user"]),
        password=os.environ.get("STARROCKS_PASSWORD", DEFAULT_SR_CONF["password"]),
        database=os.environ.get("HSGT_DATABASE", "portal_db"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=10,
    )


def _get_starrocks_conn():
    """StarRocks 连接（执行只读 SQL 导出）。"""
    return pymysql.connect(
        host=os.environ.get("STARROCKS_HOST", DEFAULT_SR_CONF["host"]),
        port=int(os.environ.get("STARROCKS_PORT", DEFAULT_SR_CONF["port"])),
        user=os.environ.get("STARROCKS_USER", DEFAULT_SR_CONF["user"]),
        password=os.environ.get("STARROCKS_PASSWORD", DEFAULT_SR_CONF["password"]),
        database=os.environ.get("STARROCKS_DATABASE", DEFAULT_SR_CONF["database"]),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=10,
        read_timeout=60,
        write_timeout=60,
    )

def _get_starrocks_meta_conn():
    """
    StarRocks 连接（专用于查询 information_schema 元数据）。
    注意：这里不要依赖 STARROCKS_DATABASE（可能不存在，例如 test_db），否则会导致连接直接失败。
    """
    return pymysql.connect(
        host=os.environ.get("STARROCKS_HOST", DEFAULT_SR_CONF["host"]),
        port=int(os.environ.get("STARROCKS_PORT", DEFAULT_SR_CONF["port"])),
        user=os.environ.get("STARROCKS_USER", DEFAULT_SR_CONF["user"]),
        password=os.environ.get("STARROCKS_PASSWORD", DEFAULT_SR_CONF["password"]),
        database="information_schema",
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        connect_timeout=10,
        read_timeout=60,
        write_timeout=60,
    )


_SQL_FORBIDDEN = re.compile(r"\b(insert|update|delete|drop|alter|create|truncate|grant|revoke|set|use|call)\b", re.I)
_SQL_HAS_LIMIT = re.compile(r"\blimit\s+\d+", re.I)
_SQL_PREVIEW_LIMIT = 500


def _leading_keyword(sql: str) -> str:
    """
    返回去掉前置注释后的首个关键字（小写）。
    仅用于语句类型判断（select/with/show/refresh）。
    """
    s = (sql or "").strip()
    if not s:
        return ""
    head_sql = s.lstrip()
    while True:
        if head_sql.startswith("--"):
            nl = head_sql.find("\n")
            if nl == -1:
                return ""
            head_sql = head_sql[nl + 1 :].lstrip()
            continue
        if head_sql.startswith("/*"):
            end = head_sql.find("*/")
            if end == -1:
                return ""
            head_sql = head_sql[end + 2 :].lstrip()
            continue
        break
    m = re.match(r"([A-Za-z_]+)", head_sql)
    return (m.group(1).lower() if m else "")


def _normalize_sql(sql: str) -> str:
    s = (sql or "").strip()
    if not s:
        raise ValueError("SQL 不能为空")
    # 禁止多语句
    if ";" in s.strip().rstrip(";"):
        raise ValueError("不支持多语句（请只提交一条 SELECT/WITH 查询）")
    # 允许前置注释（-- ... 或 /* ... */），忽略后再校验
    head_kw = _leading_keyword(s)
    if not head_kw:
        # 可能整条都是注释、或注释不完整
        head_sql = s.lstrip()
        if head_sql.startswith("/*") and "*/" not in head_sql:
            raise ValueError("SQL 注释不完整（缺少 */）")
        raise ValueError("SQL 不能为空（仅包含注释）")
    allowed_prefixes = ("select", "with", "show", "refresh")
    if head_kw not in allowed_prefixes:
        raise ValueError("仅支持 SELECT/WITH/SHOW/REFRESH（不支持多语句）")
    # 对 SELECT/WITH 做更严格的“只读”关键字校验；SHOW/REFRESH 可能包含诸如 "SHOW CREATE ..."
    # 这种字符串，否则会被误判为写入语句。
    if head_kw in ("select", "with") and _SQL_FORBIDDEN.search(s):
        raise ValueError("SQL 包含不允许的关键字（仅支持只读查询）")
    return s


def _json_safe_value(v):
    """
    Flask jsonify 默认会把 datetime/date 转成 RFC1123（GMT）字符串；
    数据中心预览希望“表里是什么就展示什么”，这里统一转成更直观的字符串：
    - date -> YYYY-MM-DD
    - datetime -> YYYY-MM-DD HH:MM:SS
    """
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dict):
        return {k: _json_safe_value(val) for k, val in v.items()}
    if isinstance(v, (list, tuple)):
        return [_json_safe_value(x) for x in v]
    return v


def _json_safe_rows(rows):
    return [_json_safe_value(r) for r in (rows or [])]


@_sql_app.route("/api/meta/tables", methods=["GET"])
def sql_to_excel_meta_tables():
    """
    查询 StarRocks 中的库/表/视图元数据，用于前端目录展示。
    仅返回只读信息：库名、表名、类型、注释。
    """
    sql = """
SELECT
  table_schema AS db_name,
  table_name,
  table_type,
  table_comment
FROM information_schema.tables
WHERE table_schema NOT IN ('information_schema', '_statistics_')
  AND table_type IN ('TABLE', 'VIEW', 'BASE TABLE')
ORDER BY table_schema, table_type, table_name
""".strip()
    try:
        conn = _get_starrocks_meta_conn()
        cur = conn.cursor()
        cur.execute(sql)
        rows = cur.fetchall() or []
        # 统一类型名：BASE TABLE 也当作 TABLE
        for r in rows:
            t = (r.get("table_type") or "").upper()
            if t == "BASE TABLE":
                r["table_type"] = "TABLE"
        cur.close()
        conn.close()
        return jsonify({"success": True, "data": rows, "count": len(rows)})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@_sql_app.route("/")
def sql_to_excel_index():
    base = (request.environ.get("SCRIPT_NAME") or "").rstrip("/")
    return render_template("sql_to_excel.html", base_path=base)


@_sql_app.route("/api/preview", methods=["POST"])
def sql_to_excel_preview():
    data = request.get_json(force=True, silent=True) or {}
    try:
        raw_sql = _normalize_sql(data.get("sql") or "")
    except ValueError as e:
        # 校验错误：返回 200 + JSON，避免上游改成 HTML 错误页
        return jsonify({"success": False, "message": str(e)})
    s = raw_sql.rstrip().rstrip(";")
    # 仅对 SELECT/WITH 追加 LIMIT 500 作为预览；SHOW/REFRESH 直接原样执行（这类语句不支持 LIMIT）。
    head_kw = _leading_keyword(s)
    if head_kw in ("select", "with") and not _SQL_HAS_LIMIT.search(s):
        preview_sql = s + f" LIMIT {_SQL_PREVIEW_LIMIT}"
    else:
        preview_sql = s
    try:
        conn = _get_starrocks_conn()
        cur = conn.cursor()
        cur.execute(preview_sql)
        rows = _json_safe_rows(cur.fetchall())
        cols = [d[0] for d in (cur.description or [])]
        cur.close()
        conn.close()
        return jsonify({"success": True, "columns": cols, "rows": rows})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@_sql_app.route("/api/export.xlsx", methods=["POST"])
def sql_to_excel_export():
    data = request.get_json(force=True, silent=True) or {}
    try:
        sql = _normalize_sql(data.get("sql") or "")
    except ValueError as e:
        # 与预览接口保持一致，返回 JSON 错误
        return jsonify({"success": False, "message": str(e)}), 400
    try:
        from tempfile import NamedTemporaryFile
        from flask import send_file, after_this_request
        from openpyxl import Workbook

        conn = _get_starrocks_conn()
        cur = conn.cursor()
        cur.execute(sql)
        cols = [d[0] for d in (cur.description or [])]

        wb = Workbook(write_only=True)
        ws = wb.create_sheet("data")
        if cols:
            ws.append(cols)

        while True:
            batch = cur.fetchmany(2000)
            if not batch:
                break
            for r in batch:
                ws.append([r.get(c) for c in cols])

        cur.close()
        conn.close()

        tmp = NamedTemporaryFile(delete=False, suffix=".xlsx")
        tmp_path = tmp.name
        tmp.close()
        wb.save(tmp_path)

        @after_this_request
        def _cleanup(resp):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass
            return resp

        return send_file(
            tmp_path,
            as_attachment=True,
            download_name="sql_export.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        return str(e), 500


@_dolphin_app.route("/api/hsgt-price-deliver", methods=["GET"])
def hsgt_price_deliver():
    """
    查询 StarRocks 中 portal_db.hsgt_price_deliver 表。

    Query 参数:
      - biz_date（必填）：业务日期，格式 yyyymmdd（如 20260211）
    """
    biz_date = (request.args.get("biz_date") or "").strip()
    if not biz_date:
        return jsonify({"success": False, "message": "biz_date 为必填参数，格式 yyyymmdd，如 20260211"}), 400
    try:
        conn = _get_hsgt_conn()
        cur = conn.cursor()
        sql = (
            "SELECT biz_date, stock_code, stock_name, last_price "
            "FROM hsgt_price_deliver "
            "WHERE biz_date = %s "
            "ORDER BY stock_code"
        )
        cur.execute(sql, (biz_date,))
        rows = cur.fetchall()
        cur.close()
        conn.close()
        return jsonify(
            {
                "success": True,
                "biz_date": biz_date,
                "count": len(rows),
                "data": rows,
            }
        )
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@_excel_app.route("/")
def excel_index():
    base = (request.environ.get("SCRIPT_NAME") or "").rstrip("/")
    return render_template("index.html", base_path=base)

@_excel_app.route("/api/impala-who")
def excel_impala_who():
    return jsonify({
        "IMPALA_USER": _excel_app.config.get("IMPALA_USER"),
        "IMPALA_RUN_AS_OS_USER": _excel_app.config.get("IMPALA_RUN_AS_OS_USER"),
        "effective_user": _excel_app.config.get("IMPALA_USER") or _excel_app.config.get("IMPALA_RUN_AS_OS_USER") or "（未配置）",
        "force_impala_shell": bool(_excel_app.config.get("IMPALA_RUN_AS_OS_USER")),
    })

@_excel_app.route("/api/upload", methods=["POST"])
def excel_upload():
    if "file" not in request.files:
        return jsonify({"success": False, "message": "未选择文件"}), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify({"success": False, "message": "未选择文件"}), 400
    if not _allowed_ext(f.filename):
        return jsonify({"success": False, "message": "仅支持 .xlsx 或 .xls 文件"}), 400
    table_name = (request.form.get("table_name") or "").strip() or Path(f.filename).stem
    table_name = re.sub(r"[^\w\u4e00-\u9fff]", "_", table_name).strip("_") or "excel_table"
    save_path = upload_dir / f.filename
    try:
        f.save(str(save_path))
        result = run_pipeline(
            excel_path=str(save_path),
            table_name=table_name,
            tmp_database=_excel_app.config["TMP_DATABASE"],
            hdfs_base_path=_excel_app.config["HDFS_BASE_PATH"],
            impala_host=_excel_app.config["IMPALA_HOST"],
            impala_port=_excel_app.config["IMPALA_PORT"],
            impala_user=_excel_app.config.get("IMPALA_USER"),
            impala_password=_excel_app.config.get("IMPALA_PASSWORD") or None,
            impala_run_as_os_user=_excel_app.config.get("IMPALA_RUN_AS_OS_USER") or None,
            hdfs_upload_as_user=_excel_app.config.get("HDFS_UPLOAD_AS_USER") or None,
            csv_encoding=_excel_app.config.get("CSV_ENCODING", "GBK"),
            table_comment=request.form.get("table_comment") or Path(f.filename).stem,
        )
        if result["ok"]:
            return jsonify({"success": True, **result})
        return jsonify({"success": False, "message": result["message"], "detail": result}), 500
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        if save_path.exists():
            try:
                save_path.unlink()
            except OSError:
                pass

# 主入口：门户样式
_main = Flask("support")
_main.config.from_object(PortalConfig)
auth_cfg = {"AUTH_COOKIE_NAME": _main.config["AUTH_COOKIE_NAME"], "SECRET_KEY": _main.config["SECRET_KEY"], "AUTH_COOKIE_MAX_AGE": _main.config["AUTH_COOKIE_MAX_AGE"]}

@_main.route("/login")
@_main.route("/logout")
def _redir():
    base = _main.config.get("PORTAL_URL", "http://127.0.0.1:5000").rstrip("/")
    p = "logout" if request.path.strip("/") == "logout" else "login"
    return redirect(f"{base}/{p}?next={request.url}")

@_main.route("/")
def _nav():
    base = (request.environ.get("SCRIPT_NAME") or "").rstrip("/")
    return redirect(f"{base}/excel_to_hive/")

app = DispatcherMiddleware(_main, {
    "/excel_to_hive": _excel_app,
    "/dolphin": _dolphin_app,
    "/sql_to_excel": _sql_app,
})
app = auth_middleware(app, auth_cfg)

if __name__ == "__main__":
    port = int(os.environ.get("SUPPORT_PORT", "5002"))
    base = os.environ.get("SUPPORT_BASE", "").rstrip("/") or ""
    print("技术支持（端口 %s）: http://127.0.0.1:%s%s" % (port, port, base or "/"))
    print("  - Excel 入库 Hive: http://127.0.0.1:%s%s/excel_to_hive/" % (port, base))
    print("  - 数据导出: http://127.0.0.1:%s%s/sql_to_excel/" % (port, base))
    from werkzeug.serving import run_simple
    # 关闭 reloader：通过 run_support.py(runpy) 启动时 reloader 子进程会报 No module named app
    use_debug = os.environ.get("SUPPORT_DEBUG", os.environ.get("PORTAL_DEBUG", "0")).lower() in ("1", "true", "yes")
    run_simple("0.0.0.0", port, app, use_debugger=use_debug, use_reloader=False, threaded=True)
